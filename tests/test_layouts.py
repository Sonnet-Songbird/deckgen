import io
import unittest
import zipfile

from pptx import Presentation

from deckgen.layouts import LAYOUTS
from deckgen.render import build
from deckgen.spec import parse_deck


def _deck(*slides: dict[str, object], theme: dict[str, object] | None = None) -> dict[str, object]:
    return {"slides": list(slides), **({"theme": theme} if theme else {})}


def _texts(slide) -> str:
    return "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)


class LayoutExamplesTest(unittest.TestCase):
    def test_모든_양식의_예제가_검증을_통과한다(self) -> None:
        for name, layout in LAYOUTS.items():
            with self.subTest(layout=name):
                self.assertEqual([], layout.validate(layout.example(), name))

    def test_모든_양식의_예제가_경고_없이_그려진다(self) -> None:
        for name, layout in LAYOUTS.items():
            with self.subTest(layout=name):
                result = build(parse_deck(_deck(layout.example())))
                self.assertEqual((), result.warnings)

    def test_양식_설명에_필드와_공통_발표자_노트가_들어_있다(self) -> None:
        for name, layout in LAYOUTS.items():
            with self.subTest(layout=name):
                names = [item["name"] for item in layout.describe()["fields"]]  # type: ignore[index]
                self.assertIn("notes", names)
                self.assertTrue(set(names) >= set(layout.example()) - {"layout"})

    def test_양식_안에_알_수_없는_키가_있으면_오류다(self) -> None:
        errors = LAYOUTS["title"].validate({"layout": "title", "title": "t", "oops": 1}, "s")
        self.assertEqual(["s.oops: 알 수 없는 항목입니다"], errors)


class RenderTest(unittest.TestCase):
    def test_슬라이드_수와_제목이_맞다(self) -> None:
        result = build(parse_deck(_deck(LAYOUTS["title"].example(), LAYOUTS["step"].example())))
        slides = list(result.presentation.slides)
        self.assertEqual(2, len(slides))
        self.assertEqual("발표 제목", slides[0].shapes.title.text_frame.text)
        self.assertEqual("STEP 1  서비스 재시작", slides[1].shapes.title.text_frame.text)

    def test_단계_양식에_명령과_성공_실패_문구가_들어간다(self) -> None:
        slide = list(build(parse_deck(_deck(LAYOUTS["step"].example()))).presentation.slides)[0]
        text = _texts(slide)
        self.assertIn("Restart-Service <서비스이름>", text)
        self.assertIn("화면이 정상으로 열리면 끝.", text)
        self.assertIn("같은 오류가 나오면 다음 단계로.", text)

    def test_바꿀_곳_표시는_붉은_굵은_글씨가_된다(self) -> None:
        slide = list(build(parse_deck(_deck(LAYOUTS["step"].example()))).presentation.slides)[0]
        slot_runs = [
            run
            for shape in slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
            if run.text == "<서비스이름>"
        ]
        self.assertEqual(1, len(slot_runs))
        self.assertTrue(slot_runs[0].font.bold)
        self.assertEqual("C0392B", str(slot_runs[0].font.color.rgb))

    def test_밝은_양식에만_쪽_번호_필드가_있다(self) -> None:
        result = build(parse_deck(_deck(LAYOUTS["title"].example(), LAYOUTS["step"].example(), LAYOUTS["closing"].example())))
        xml = [slide._element.xml for slide in result.presentation.slides]
        flags = ['type="slidenum"' in item for item in xml]
        self.assertEqual([False, True, False], flags)

    def test_발표자_노트가_저장된다(self) -> None:
        slide_data = {**LAYOUTS["title"].example(), "notes": "발표 때 말할 내용"}
        slide = list(build(parse_deck(_deck(slide_data))).presentation.slides)[0]
        self.assertEqual("발표 때 말할 내용", slide.notes_slide.notes_text_frame.text)

    def test_테마_색이_적용된다(self) -> None:
        result = build(parse_deck(_deck(LAYOUTS["checklist"].example(), theme={"palette": {"ok": "123456"}})))
        self.assertIn('srgbClr val="123456"', list(result.presentation.slides)[0]._element.xml)

    def test_저장한_파일을_다시_열_수_있다(self) -> None:
        data = build(parse_deck(_deck(*(layout.example() for layout in LAYOUTS.values())))).to_bytes()
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(data)))
        self.assertEqual(len(LAYOUTS), len(Presentation(io.BytesIO(data)).slides))

    def test_흐름도는_시작_번호와_마지막_강조를_따른다(self) -> None:
        flow = {**LAYOUTS["flow"].example(), "start": 1, "highlight_last": True, "tag": "단계"}
        slide = list(build(parse_deck(_deck(flow))).presentation.slides)[0]
        self.assertIn("단계\n1\n확인", _texts(slide))
        self.assertIn("단계\n3\n검증", _texts(slide))


class OverflowWarningTest(unittest.TestCase):
    def test_성공_실패_문구가_너무_길면_경고한다(self) -> None:
        step = {**LAYOUTS["step"].example(), "ng": "가" * 100}
        warnings = build(parse_deck(_deck(step))).warnings
        self.assertTrue(any("성공/실패 문구가 길어" in item.message for item in warnings))
        self.assertEqual(1, warnings[0].slide)

    def test_긴_명령_줄은_줄바꿈_경고를_낸다(self) -> None:
        step = {**LAYOUTS["step"].example(), "code": ["x" * 94]}
        warnings = build(parse_deck(_deck(step))).warnings
        self.assertTrue(any("줄바꿈됩니다" in item.message for item in warnings))

    def test_본문이_성공_실패_상자와_겹치면_경고한다(self) -> None:
        step = {**LAYOUTS["step"].example(), "code": ["a"] * 8, "extra": "추가 안내"}
        warnings = build(parse_deck(_deck(step))).warnings
        self.assertTrue(any("겹칩니다" in item.message for item in warnings))


if __name__ == "__main__":
    unittest.main()
