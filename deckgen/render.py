"""DeckSpec 으로 pptx 를 만든다."""

import io
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.util import Inches

from deckgen.painter import SLIDE_H, SLIDE_W, BuildWarning, SlidePainter
from deckgen.spec import DeckSpec
from deckgen.theme import rgb

_TITLE_ONLY_LAYOUT = 5


@dataclass(frozen=True)
class BuildResult:
    presentation: PresentationType
    warnings: tuple[BuildWarning, ...]

    def save(self, path: Path) -> None:
        self.presentation.save(str(path))

    def to_bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.presentation.save(buffer)
        return buffer.getvalue()


def build(deck: DeckSpec) -> BuildResult:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    palette = deck.theme.palette
    warnings: list[BuildWarning] = []
    for index, spec in enumerate(deck.slides, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_ONLY_LAYOUT])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb({True: palette.primary, False: palette.paper}[spec.layout.dark])
        painter = SlidePainter(slide, deck.theme, index, spec.layout.name, warnings)
        spec.layout.render(painter, spec.data)
        if not spec.layout.dark:
            painter.slide_number()
        slide.notes_slide.notes_text_frame.text = str(spec.data.get("notes", ""))
    presentation.core_properties.title = str(deck.slides[0].data.get("title", ""))
    return BuildResult(presentation, tuple(warnings))
