"""슬라이드 한 장에 도형과 글을 그리는 도구. 양식은 이 도구만 호출한다."""

import re
from collections.abc import Sequence
from dataclasses import dataclass

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.util import Inches, Pt

from deckgen.metrics import count_lines, lines_height_in
from deckgen.theme import Theme, blend, rgb

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN_X = 0.7
CONTENT_W = SLIDE_W - 2 * MARGIN_X
RESULT_ROW_Y = 5.75
RESULT_ROW_MIN_H = 1.05
RESULT_ROW_MAX_H = 1.25
FIT_TOLERANCE_IN = 0.05

_ANCHORS = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_ALIGNS = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_SLOT = re.compile(r"(<[^>]+>)")
_SLIDE_NUMBER_FIELD_ID = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"


@dataclass(frozen=True)
class BuildWarning:
    slide: int
    layout: str
    message: str


@dataclass(frozen=True)
class Run:
    text: str
    size: float
    color: str
    bold: bool = False
    mono: bool = False
    symbol: bool = False


@dataclass(frozen=True)
class Para:
    runs: tuple[Run, ...]
    align: str = "left"
    space_after: float = 0.0

    @property
    def size(self) -> float:
        return max(run.size for run in self.runs)

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.runs)


def slot_runs(text: str, size: float, color: str, slot_color: str, bold: bool = False, mono: bool = False) -> tuple[Run, ...]:
    """<...> 로 감싼 부분은 바꿔 써야 하는 자리라서 붉은 굵은 고정폭 글씨로 만든다."""
    parts = [part for part in _SLOT.split(text) if part]
    return tuple(
        Run(part, size, slot_color if _SLOT.fullmatch(part) else color, bold or bool(_SLOT.fullmatch(part)), mono or bool(_SLOT.fullmatch(part)))
        for part in parts
    )


class SlidePainter:
    def __init__(self, slide: Slide, theme: Theme, index: int, layout_name: str, warnings: list[BuildWarning]) -> None:
        self.slide = slide
        self.theme = theme
        self.index = index
        self.layout_name = layout_name
        self.warnings = warnings

    # ---- 경고 -------------------------------------------------------------
    def warn(self, message: str) -> None:
        self.warnings.append(BuildWarning(self.index, self.layout_name, message))

    def warn_if_below(self, bottom: float, limit: float, message: str) -> None:
        """본문 아래쪽 끝(bottom)이 한계(limit)를 넘으면 경고한다."""
        self.warnings.extend([BuildWarning(self.index, self.layout_name, message)] * int(bottom > limit + FIT_TOLERANCE_IN))

    # ---- 도형 -------------------------------------------------------------
    def _shape(self, kind: MSO_SHAPE, x: float, y: float, w: float, h: float, fill: str, name: str) -> Shape:
        shape = self.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = name
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.fill.background()
        shape._element.remove(shape._element.find(qn("p:style")))
        return shape

    def card(self, x: float, y: float, w: float, h: float, fill: str, name: str, radius: float = 0.08) -> None:
        shape = self._shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, name)
        shape.adjustments[0] = min(0.5, radius / min(w, h))

    def circle(self, x: float, y: float, diameter: float, fill: str, name: str) -> None:
        self._shape(MSO_SHAPE.OVAL, x, y, diameter, diameter, fill, name)

    # ---- 글 ---------------------------------------------------------------
    def text(self, paragraphs: Sequence[Para], x: float, y: float, w: float, h: float, name: str, anchor: str = "top") -> None:
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        self._write(box.text_frame, paragraphs, w, h, anchor)

    def _write(self, frame, paragraphs: Sequence[Para], w: float, h: float, anchor: str) -> None:
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = _ANCHORS[anchor]
        needed = 0.0
        for index, para in enumerate(paragraphs):
            target = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            target.alignment = _ALIGNS[para.align]
            target.space_after = Pt(para.space_after)
            for run in para.runs:
                self._add_run(target, run)
            lines = max(1, sum(count_lines(line, para.size, w, para.runs[0].mono, para.runs[0].bold) for line in para.text.split("\n")))
            needed += lines_height_in(lines, para.size) + para.space_after / 72.0
        overflow = needed > h + FIT_TOLERANCE_IN
        self.warnings.extend([BuildWarning(self.index, self.layout_name, f"글이 상자를 넘칠 수 있습니다: {paragraphs[0].text[:24]!r} (필요 {needed:.2f}in > 상자 {h:.2f}in)")] * int(overflow))

    def _add_run(self, paragraph, run: Run) -> None:
        added = paragraph.add_run()
        added.text = run.text
        font = added.font
        font.size = Pt(run.size)
        font.bold = run.bold
        font.color.rgb = rgb(run.color)
        font_name = {(False, False): self.theme.font, (True, False): self.theme.code_font}.get((run.mono, run.symbol), self.theme.symbol_font)
        font.name = font_name
        latin = added._r.get_or_add_rPr().find(qn("a:latin"))
        east_asian = etree.Element(qn("a:ea"))
        east_asian.set("typeface", font_name)
        latin.addnext(east_asian)

    def title(self, text: str, x: float, y: float, w: float, h: float, size: float, color: str, anchor: str = "middle") -> None:
        placeholder = self.slide.shapes.title
        placeholder.left, placeholder.top, placeholder.width, placeholder.height = Inches(x), Inches(y), Inches(w), Inches(h)
        placeholder.name = "제목"
        self._write(placeholder.text_frame, [Para((Run(text, size, color, bold=True),))], w, h, anchor)

    # ---- 자주 쓰는 조각 -----------------------------------------------------
    def badge(self, label: str, tone: str) -> None:
        fill = {"normal": self.theme.palette.accent, "warn": self.theme.palette.warn}[tone]
        self.circle(MARGIN_X, 0.4, 0.95, fill, "단계 번호 배경")
        self.text([Para((Run(label, 28, self.theme.palette.paper, bold=True),), align="center")], MARGIN_X, 0.4, 0.95, 0.95, "단계 번호", anchor="middle")

    def numbered_circle(self, x: float, y: float, diameter: float, label: str, fill: str, text_color: str, size: float, name: str) -> None:
        self.circle(x, y, diameter, fill, f"{name} 배경")
        self.text([Para((Run(label, size, text_color, bold=True),), align="center")], x, y, diameter, diameter, name, anchor="middle")

    def symbol_circle(self, x: float, y: float, diameter: float, symbol: str, fill: str, size: float, name: str) -> None:
        self.circle(x, y, diameter, fill, f"{name} 배경")
        self.text([Para((Run(symbol, size, self.theme.palette.paper, bold=True, symbol=True),), align="center")], x, y, diameter, diameter, name, anchor="middle")

    def lead(self, text: str, y: float = 1.6) -> float:
        """안내 문장. 한 줄이면 0.5in, 두 줄이면 0.85in 를 쓰고 그 높이를 돌려준다."""
        lines = count_lines(text, 18, CONTENT_W)
        height = {1: 0.5}.get(lines, 0.85)
        self.text([Para((Run(text, 18, self.theme.palette.ink),),)], MARGIN_X, y, CONTENT_W, height, "할 일", anchor="middle")
        return height

    def code_block(self, lines: Sequence[str], y: float) -> float:
        """코드 상자. 아래쪽 y 를 돌려준다."""
        palette = self.theme.palette
        height = len(lines) * 0.36 + 0.5
        self.card(MARGIN_X, y, CONTENT_W, height, palette.surface, "명령 배경")
        paragraphs = [Para(slot_runs(line, 15, palette.ink, palette.danger, mono=True)) for line in lines]
        self.text(paragraphs, MARGIN_X + 0.3, y + 0.15, CONTENT_W - 0.6, height - 0.3, "명령", anchor="middle")
        widest = max(count_lines(line, 15, CONTENT_W - 0.6, mono=True) for line in lines)
        self.warnings.extend([BuildWarning(self.index, self.layout_name, "코드 줄이 길어 상자 안에서 줄바꿈됩니다. 줄을 나누세요")] * int(widest > 1))
        return y + height

    def result_row(self, ok: str, ng: str, y: float = RESULT_ROW_Y) -> None:
        """성공/실패 상자 한 쌍. 글이 길면 높이를 늘리고 한계를 넘을 때만 경고한다."""
        palette = self.theme.palette
        gap = 0.3
        w = (CONTENT_W - gap) / 2
        inner_w = w - 0.5
        needed = max(lines_height_in(1, 18) + lines_height_in(count_lines(body, 16, inner_w), 16) for body in (ok, ng)) + 0.2
        height = min(RESULT_ROW_MAX_H, max(RESULT_ROW_MIN_H, needed))
        self.warn_if_below(needed, RESULT_ROW_MAX_H, "성공/실패 문구가 길어 상자를 넘칩니다. 문구를 줄이세요")
        for index, (label, body, color) in enumerate((("성공", ok, palette.ok), ("실패", ng, palette.warn))):
            x = MARGIN_X + index * (w + gap)
            self.card(x, y, w, height, blend(color, palette.paper, 0.14), f"{label} 상자")
            self.text(
                [Para((Run(label, 18, color, bold=True),)), Para((Run(body, 16, palette.ink),))],
                x + 0.25, y, inner_w, height, f"{label} 문구", anchor="middle",
            )

    def slide_number(self) -> None:
        """슬라이드 번호 필드. 순서를 바꾸면 PowerPoint 가 번호를 다시 매긴다."""
        box = self.slide.shapes.add_textbox(Inches(SLIDE_W - 1.2), Inches(7.0), Inches(0.6), Inches(0.3))
        box.name = "쪽 번호"
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        field = etree.SubElement(paragraph._p, qn("a:fld"))
        field.set("id", _SLIDE_NUMBER_FIELD_ID)
        field.set("type", "slidenum")
        properties = etree.SubElement(field, qn("a:rPr"))
        properties.set("lang", "ko-KR")
        properties.set("sz", "1200")
        fill = etree.SubElement(properties, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", self.theme.palette.muted)
        etree.SubElement(properties, qn("a:latin")).set("typeface", self.theme.font)
        etree.SubElement(field, qn("a:t")).text = "‹#›"
